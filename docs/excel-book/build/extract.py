"""Extract recalculated project values into mock/projdata.js for the illustrations."""
import json, datetime as dt
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
P = '../projets/'
def fmt(v, nf):
    if v is None: return ''
    if isinstance(v, (dt.datetime, dt.date)): return v.strftime('%d/%m/%Y')
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        if '%' in nf: return (f'{v*100:,.1f}'.replace(',', ' ').replace('.', ',')) + ' %'
        if '0.00' in nf: return f'{v:,.2f}'.replace(',', ' ').replace('.', ',')
        if nf in ('0',): return f'{v:,.0f}'.replace(',', ' ')
        return (f'{v:,.2f}' if isinstance(v, float) and v != int(v) else f'{int(v):,}').replace(',', ' ').replace('.', ',')
    return str(v)
STATE = {'Dépassement': 'background:#F8CBAD;color:#9C0006;font-weight:700', 'Rupture': 'background:#F8CBAD;color:#9C0006;font-weight:700', 'Ajourné': 'background:#F8CBAD;color:#9C0006;font-weight:700',
         'À commander': 'background:#FFEB9C;color:#9C5700;font-weight:700', 'OK': 'background:#C6EFCE;color:#006100', 'Admis': 'background:#C6EFCE;color:#006100;font-weight:700', 'A': 'background:#F8CBAD;color:#9C0006;font-weight:700', 'C': 'background:#DDEBF7'}
def grab(f, sh, rng, state_cols=()):
    wb = load_workbook(P + f, data_only=True); ws = wb[sh]
    a, b = rng.split(':'); c1, r1 = CI(''.join(x for x in a if x.isalpha())), int(''.join(x for x in a if x.isdigit())); c2, r2 = CI(''.join(x for x in b if x.isalpha())), int(''.join(x for x in b if x.isdigit()))
    merged = {str(m).split(':')[0]: str(m).split(':')[1] for m in ws.merged_cells.ranges}
    covered = set()
    for m in ws.merged_cells.ranges:
        for row in ws.iter_rows(min_row=m.min_row, max_row=m.max_row, min_col=m.min_col, max_col=m.max_col):
            for x in row:
                if x.coordinate != str(m).split(':')[0]: covered.add(x.coordinate)
    cells, widths = {}, {}
    for c in range(c1, c2 + 1):
        widths[L(c)] = round((ws.column_dimensions[L(c)].width or 8.43) * 7.2 + 6)
        for r in range(r1, r2 + 1):
            x = ws.cell(r, c)
            if x.coordinate in covered: continue
            v = fmt(x.value, x.number_format or 'General'); st = []
            fill = x.fill.fgColor.rgb if x.fill and x.fill.fill_type == 'solid' else None
            fill = fill[-6:] if isinstance(fill, str) else None
            if fill: st.append(f'background:#{fill}')
            if x.font and x.font.bold: st.append('font-weight:700')
            if x.font and x.font.color is not None and isinstance(x.font.color.rgb, str): st.append(f'color:#{x.font.color.rgb[-6:]}')
            if x.font and x.font.italic: st.append('font-style:italic')
            if x.font and x.font.size and x.font.size >= 14: st.append(f'font-size:{x.font.size+3}px')
            if x.alignment and x.alignment.horizontal == 'center': st.append('justify-content:center')
            if x.alignment and x.alignment.horizontal == 'right': st.append('justify-content:flex-end')
            if v in STATE and (L(c) in state_cols): st.append(STATE[v])
            if x.border and x.border.left and x.border.left.style: st.append('border:1px solid #bfbfbf')
            if not v and not st: continue
            d = {'v': v, 's': ';'.join(st)}
            if x.coordinate in merged: d['span'] = merged[x.coordinate]
            cells[x.coordinate] = d
    return {'cells': cells, 'w': widths}
D = {
 'X26': grab('Projet1-Budget-personnel.xlsx', 'Budget', 'A1:F26', 'F'),
 'X27': grab('Projet2-Facture-automatique.xlsx', 'Facture', 'A1:F27'),
 'X28': grab('Projet3-Gestion-de-stock.xlsx', 'Articles', 'A1:K17', 'K'),
 'X29': grab('Projet4-Releve-de-notes.xlsx', 'Relevé', 'A1:K22', 'K'),
 'X30': grab('Projet5-Tableau-de-bord-ventes.xlsx', 'Tableau de bord', 'A1:I28'),
 'X31': grab('Projet6-Presence-et-paie.xlsx', 'Présence', 'A1:AL13', 'CDEFGHIJKLMNOPQRSTUVWXYZ'),
}
D['X31']['cells'] = {k: v for k, v in D['X31']['cells'].items()}
for k in list(D['X31']['cells']):
    col = ''.join(ch for ch in k if ch.isalpha())
    if col not in ('A', 'B', 'AH', 'AI', 'AJ', 'AK', 'AL') and len(col) <= 2:
        v = D['X31']['cells'][k]
        if v['v'] in STATE: v['s'] += ';' + STATE[v['v']]
open('mock/projdata.js', 'w').write('const PD = ' + json.dumps(D, ensure_ascii=False) + ';\n')
print({k: len(v['cells']) for k, v in D.items()})
