"""Workbook that checks the formulas quoted in the Excel book.
Each row: the French formula as printed, the same formula live (English names, as stored by Excel),
the expected result, the obtained result and OK / À vérifier. Rows marked MANUEL must be tested by hand in Excel."""
import datetime as dt, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'Tests-formules-Excel.xlsx')
wb = Workbook(); t = wb.active; t.title = 'Tests'; d = wb.create_sheet('Data')
A = Font(name='Arial'); H = Font(name='Arial', bold=True, color='FFFFFF'); G = PatternFill('solid', fgColor='217346')

# --- data
for i, v in enumerate('ABC'): d.cell(1 + i, 1, v)
prods = [('P001', 'Cahier', 120, 80), ('P002', 'Stylo', 30, 300), ('P003', 'Classeur', 250, 45), ('P004', 'Ramette A4', 750, 20), ('P005', 'Agenda', 400, 15), ('P006', 'Calculatrice', 1200, 9)]
for j, h in enumerate(['Code', 'Produit', 'Prix', 'Stock']): d.cell(1, 4 + j, h)
for i, r in enumerate(prods):
    for j, v in enumerate(r): d.cell(2 + i, 4 + j, v)
emps = [('Ali', 'Commercial', 120000, 135000), ('Fatima', 'Commercial', 98000, 87000), ('Hamza', 'Marketing', 76000, 99000), ('Salma', 'Commercial', 143000, 151000), ('Youssef', 'Marketing', 54000, 61000), ('Asma', 'Commercial', 110000, 104000)]
for j, h in enumerate(['Employé', 'Service', 'T1', 'T2']): d.cell(1, 9 + j, h)
for i, r in enumerate(emps):
    for j, v in enumerate(r): d.cell(2 + i, 9 + j, v)
d['M1'], d['N1'] = 'Élève', 'Moyenne'
for i, (n, m) in enumerate([('Amine', 16.96), ('Sara', 9.4), ('Yacine', 12.75), ('Lina', 8.2), ('Omar', 14.1)]): d.cell(2 + i, 13, n); d.cell(2 + i, 14, m)
d['P1'] = 'Dates'
for i, v in enumerate([dt.date(1998, 3, 14), dt.date(2026, 10, 1), dt.date(2026, 10, 31), dt.date(2026, 10, 5)]): d.cell(2 + i, 16, v).number_format = 'dd/mm/yyyy'
d['R1'] = 'Textes'
for i, v in enumerate(['  sami   EXEMPLE ', 'Sami Exemple', 'P004-ALG', 'F-2026-0042', 'Sami', 'Exemple']): d.cell(2 + i, 18, v)
d['T1'] = 'Notes / coefficients'
for j, (n, c) in enumerate(zip([17.5, 17.75, 16.75, 18.25, 13.75], [4, 3, 3, 2, 2])): d.cell(2, 20 + j, n); d.cell(3, 20 + j, c)

# --- tests: (chapitre, formule FR imprimée, formule EN vivante ou None, attendu)
D = 'Data!'
T = [
 (5, '=(2+3)*4', '=(2+3)*4', 20), (5, '=2+3*4', '=2+3*4', 14), (5, '=A2&" "&B2', f'={D}R6&" "&{D}R7', 'Sami Exemple'),
 (6, '=SOMME(C5:C10)', f'=SUM({D}K2:K7)', 601000), (6, '=MOYENNE(D5:D10)', f'=ROUND(AVERAGE({D}L2:L7),2)', 106166.67), (6, '=MAX(D5:D10)', f'=MAX({D}L2:L7)', 151000),
 (6, '=NB(…)', f'=COUNT({D}I2:L7)', 12), (6, '=NBVAL(…)', f'=COUNTA({D}I2:I7)', 6), (6, '=ARRONDI(B2;2)', '=ROUND(2.456,2)', 2.46),
 (7, '=SI(B2>=10;"Admis";"Ajourné")', f'=IF({D}N2>=10,"Admis","Ajourné")', 'Admis'), (7, '=SI(B3>=10;"Admis";"Ajourné")', f'=IF({D}N3>=10,"Admis","Ajourné")', 'Ajourné'),
 (7, '=SI(ET(B2>=10;C2>=10);"Validé";"Non validé")', f'=IF(AND({D}N2>=10,{D}N4>=10),"Validé","Non validé")', 'Validé'),
 (7, '=SI(OU(K5="Rupture";K5="À commander");"Commander";"")', '=IF(OR("À commander"="Rupture","À commander"="À commander"),"Commander","")', 'Commander'),
 (7, 'SI imbriqués (H7 = 12,75)', f'=IF({D}N4>=16,"Très bien",IF({D}N4>=14,"Bien",IF({D}N4>=12,"Assez bien",IF({D}N4>=10,"Passable",""))))', 'Assez bien'),
 (7, '=SI.CONDITIONS(…) (H7 = 14,1)', f'=_xlfn.IFS({D}N6>=16,"Très bien",{D}N6>=14,"Bien",{D}N6>=12,"Assez bien",{D}N6>=10,"Passable",TRUE,"")', 'Bien'),
 (7, '=NB.SI(B5:B10;"Commercial")', f'=COUNTIF({D}J2:J7,"Commercial")', 4), (7, '=SOMME.SI(B5:B10;"Marketing";C5:C10)', f'=SUMIF({D}J2:J7,"Marketing",{D}K2:K7)', 130000),
 (7, '=SOMME.SI.ENS(…;"Commercial";…;">100000")', f'=SUMIFS({D}K2:K7,{D}J2:J7,"Commercial",{D}L2:L7,">100000")', 373000), (7, '=NB.SI.ENS(…;"Commercial";…;">100000")', f'=COUNTIFS({D}J2:J7,"Commercial",{D}L2:L7,">100000")', 3),
 (7, '=MOYENNE.SI(…;"Marketing";…)', f'=AVERAGEIF({D}J2:J7,"Marketing",{D}K2:K7)', 65000), (7, '=SIERREUR(B2/C2;0)', '=IFERROR(1/0,0)', 0),
 (13, 'SOMME.SI.ENS avec "*" (Toutes)', f'=SUMIFS({D}K2:K7,{D}J2:J7,"*")', 601000),
 (8, '=RECHERCHEV(B13;$A$5:$D$10;2;FAUX)', f'=VLOOKUP("P004",{D}D2:G7,2,FALSE)', 'Ramette A4'), (8, '=RECHERCHEV(…;3;FAUX)', f'=VLOOKUP("P004",{D}D2:G7,3,FALSE)', 750),
 (8, '=EQUIV("C";A1:A3;0)', f'=MATCH("C",{D}A1:A3,0)', 3), (8, '=INDEX(B5:B10;4)', f'=INDEX({D}E2:E7,4)', 'Ramette A4'),
 (8, '=INDEX(B5:B10;EQUIV(B13;A5:A10;0)) (P006)', f'=INDEX({D}E2:E7,MATCH("P006",{D}D2:D7,0))', 'Calculatrice'),
 (8, '=INDEX(B7:B22;EQUIV(MAX(H7:H22);H7:H22;0))', f'=INDEX({D}M2:M6,MATCH(MAX({D}N2:N6),{D}N2:N6,0))', 'Amine'),
 (8, '=RECHERCHEX(B13;A5:A10;B5:B10;"Introuvable")', None, 'Ramette A4 (MANUEL : Microsoft 365 / Excel 2021+)'),
 (9, '=SUPPRESPACE("  sami   EXEMPLE ")', f'=TRIM({D}R2)', 'sami EXEMPLE'), (9, '=NOMPROPRE(SUPPRESPACE(A5))', f'=PROPER(TRIM({D}R2))', 'Sami Exemple'),
 (9, '=MAJUSCULE("p004")', '=UPPER("p004")', 'P004'), (9, '=GAUCHE("P004-ALG";4)', f'=LEFT({D}R4,4)', 'P004'), (9, '=STXT("F-2026-0042";3;4)', f'=MID({D}R5,3,4)', '2026'),
 (9, '=NBCAR("Excel")', '=LEN("Excel")', 5), (9, '=TROUVE(" ";"Sami Exemple")', f'=FIND(" ",{D}R3)', 5), (9, '=GAUCHE(B5;TROUVE(" ";B5)-1)', f'=LEFT({D}R3,FIND(" ",{D}R3)-1)', 'Sami'),
 (9, '=MAJUSCULE(GAUCHE(C5;3))', f'=UPPER(LEFT({D}R6,3))', 'SAM'), (9, '=DATEDIF(A2;date;"y")', f'=DATEDIF({D}P2,{D}P5,"y")', 28),
 (9, '=NB.JOURS.OUVRES(début;fin) (octobre 2026)', f'=NETWORKDAYS({D}P3,{D}P4)', 22), (9, '=NB.JOURS.OUVRES.INTL(début;fin;7)', f'=NETWORKDAYS.INTL({D}P3,{D}P4,7)', 21),
 (9, '=FIN.MOIS(A2;0) ← 31/10/2026', f'=EOMONTH({D}P5,0)=DATE(2026,10,31)', True), (9, '=DATE(2026;10;5)', f'=DATE(2026,10,5)={D}P5', True),
 (9, '=ANNEE(A2) / MOIS / JOUR', f'=YEAR({D}P5)*10000+MONTH({D}P5)*100+DAY({D}P5)', 20261005),
 (9, '=TEXTE(A2;"jjjj")', None, 'lundi (MANUEL : codes français jj / aaaa)'), (9, '=TEXTE(A2;"mmmm aaaa")', None, 'octobre 2026 (MANUEL)'),
 (10, '=SOMME(Ventes[Montant])', None, 'MANUEL : créer le tableau « Ventes » (Ctrl+T) dans Projet5, résultat 20 744 300'),
 (13, '=ARRONDI(SOMMEPROD(C7:G7;$C$4:$G$4)/SOMME($C$4:$G$4);2)', f'=ROUND(SUMPRODUCT({D}T2:X2,{D}T3:X3)/SUM({D}T3:X3),2)', 16.96),
 (13, '=RANG(H7;$H$7:$H$22;0)', f'=RANK({D}N2,{D}N2:N6,0)', 1),
]
heads = ['N°', 'Chapitre', 'Formule (FR, telle qu’imprimée)', 'Formule testée (EN)', 'Résultat attendu', 'Résultat obtenu', 'Vérification']
for j, h in enumerate(heads):
    c = t.cell(1, 1 + j, h); c.font = H; c.fill = G; c.alignment = Alignment(horizontal='center')
for i, (ch, fr, en, exp) in enumerate(T):
    r = 2 + i
    t.cell(r, 1, i + 1); t.cell(r, 2, ch); x = t.cell(r, 3, fr); x.data_type = 's'
    t.cell(r, 4, en[1:] if en else '—'); t.cell(r, 5, exp)
    if en:
        t.cell(r, 6, en)
        t.cell(r, 7, f'=IF(F{r}=E{r},"OK","À vérifier")')
    else:
        t.cell(r, 6, '—'); t.cell(r, 7, 'MANUEL')
    for c in range(1, 8): t.cell(r, c).font = A
t.conditional_formatting.add(f'G2:G{len(T)+1}', CellIsRule(operator='equal', formula=['"OK"'], fill=PatternFill('solid', fgColor='C6EFCE')))
t.conditional_formatting.add(f'G2:G{len(T)+1}', CellIsRule(operator='equal', formula=['"MANUEL"'], fill=PatternFill('solid', fgColor='FFEB9C')))
for col, w in zip('ABCDEFG', [5, 9, 52, 58, 30, 18, 13]): t.column_dimensions[col].width = w
t.freeze_panes = 'A2'
r = len(T) + 3
t.cell(r, 3, 'Résumé :').font = Font(name='Arial', bold=True)
t.cell(r, 4, f'=COUNTIF(G2:G{len(T)+1},"OK")&" OK · "&COUNTIF(G2:G{len(T)+1},"À vérifier")&" à vérifier · "&COUNTIF(G2:G{len(T)+1},"MANUEL")&" manuels"')
wb.calculation.fullCalcOnLoad = True
wb.save(OUT); print(len(T), 'tests')
